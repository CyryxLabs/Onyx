"""Governed, source-bound multi-format document ingestion for Phase 7.

The ingestor is deliberately read-only. It receives immutable bytes, verifies
an accepted source, extracts granular cited units, and returns reviewable
memory candidates. Documents and observations are always untrusted evidence
and never gain runtime instruction authority.
"""

from __future__ import annotations

import ast
import csv
import difflib
import hashlib
import hmac
import io
import json
import os
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from typing import Final, NoReturn

from core.phase7_approved_sources_v1 import (
    ApprovedSourceRecordV1,
    ApprovedSourceRegistryV1,
    ApprovedSourceV1Denied,
)
from core.phase7_layered_memory_v1 import LayeredMemorySpecV1

FEATURE_FLAG: Final = "ONYX_PHASE7_DOCUMENT_INGESTION_V1"
ENABLED_VALUE: Final = "true"
SCHEMA: Final = "OnyxGovernedDocumentIngestion.v1"
MAX_DOCUMENT_BYTES: Final = 32 * 1024 * 1024
MAX_EXPANDED_BYTES: Final = 128 * 1024 * 1024
MAX_ARCHIVE_ENTRIES: Final = 10_000
MAX_UNITS: Final = 25_000
MAX_UNIT_CHARS: Final = 20_000
MAX_TOTAL_TEXT_CHARS: Final = 4_000_000
MAX_OCR_BLOCKS: Final = 10_000

FORMATS: Final = {
    "application/pdf": ("pdf", {".pdf"}),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
        "docx",
        {".docx"},
    ),
    "text/markdown": ("markdown", {".md", ".markdown"}),
    "text/plain": ("text", {".txt", ".rst", ".log"}),
    "text/csv": ("spreadsheet", {".csv"}),
    "text/tab-separated-values": ("spreadsheet", {".tsv"}),
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": (
        "spreadsheet",
        {".xlsx", ".xlsm"},
    ),
    "application/vnd.ms-excel": ("spreadsheet", {".xls"}),
    "application/vnd.oasis.opendocument.spreadsheet": (
        "spreadsheet",
        {".ods"},
    ),
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": (
        "presentation",
        {".pptx"},
    ),
    "image/png": ("image", {".png"}),
    "image/jpeg": ("image", {".jpg", ".jpeg"}),
    "image/webp": ("image", {".webp"}),
    "image/tiff": ("image", {".tif", ".tiff"}),
    "image/bmp": ("image", {".bmp"}),
    "text/x-python": ("code", {".py"}),
    "text/javascript": ("code", {".js", ".mjs", ".cjs"}),
    "text/typescript": ("code", {".ts", ".tsx"}),
    "text/x-source-code": (
        "code",
        {
            ".c",
            ".cc",
            ".cpp",
            ".cs",
            ".go",
            ".java",
            ".kt",
            ".php",
            ".rb",
            ".rs",
            ".sh",
            ".sql",
            ".swift",
            ".yaml",
            ".yml",
            ".toml",
        },
    ),
}

SOURCE_ENTRY_ROOTS: Final = (
    (
        "docs/onyx/checkpoints/phase7-approved-sources-company-graph-v1/manifest.json",
        "fa5fea91caf6914e559cbc22420841b602eeace0153653dc57591fb0070f23f3",
    ),
    (
        "docs/onyx/acceptance/VE-P7-APPROVED-SOURCES-COMPANY-GRAPH-V1-E6-001.md",
        "b5336b1484c1225b58454dd39cfeb6fad51662fd46b7658bfbb3572389222650",
    ),
    (
        "docs/onyx/acceptance/VE-P7-APPROVED-SOURCES-COMPANY-GRAPH-V1-E6-001.manifest.json",
        "27522c63449d625c97b500ac8ed9048aeddb4cf1d20f77a96bd5431cd21cd14d",
    ),
    (
        "docs/onyx/VE-ACCEPTANCE-P7-APPROVED-SOURCES-COMPANY-GRAPH-V1-E6-001.sha256",
        "d5d17c9fdceb76b43405afc8b54129f1e6c9af1e35e67f140cb25aeafb674f85",
    ),
)
LAYERED_MEMORY_ENTRY_ROOTS: Final = (
    (
        "docs/onyx/checkpoints/phase7-layered-memory-v1/manifest.json",
        "2818c6ffe66681436286d89f557c57cd23c35c29ee7f29d53d7eb2bf9d8ec95d",
    ),
    (
        "docs/onyx/acceptance/VE-P7-LAYERED-MEMORY-V1-E6-001.md",
        "594af1ea18644aa3f2ab951ecf2204bde01649399cb8a2c8ca793308b87def5a",
    ),
    (
        "docs/onyx/acceptance/VE-P7-LAYERED-MEMORY-V1-E6-001.manifest.json",
        "5fb4ea8b702967a4d9861215783f54b79ab156637870c195d32ee7bb4ccae0d6",
    ),
    (
        "docs/onyx/VE-ACCEPTANCE-P7-LAYERED-MEMORY-V1-E6-001.sha256",
        "b4c14e068728d685f22b9049d0f2776188072451eb1dde850081238f41dd3fd7",
    ),
)

_SAFE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:@/-]{0,159}$")
_FILENAME = re.compile(r"^[^\\/:*?\"<>|\x00-\x1f]{1,240}$")
_HEX64 = re.compile(r"^[0-9a-f]{64}$")
_REQ = re.compile(
    r"(?:\b(?:must|shall|required|requirement|needs? to|deve|deverá|precisa)\b|(?:REQ|FR|NFR)-\d+)",
    re.IGNORECASE,
)
_DECISION = re.compile(
    r"\b(?:decision|decided|approved|rejected|we will|decisão|decidido|aprovado|rejeitado)\b",
    re.IGNORECASE,
)
_POISON = (
    (
        "instruction_override",
        re.compile(r"\bignore (?:all |the )?(?:previous|prior) instructions?\b", re.I),
    ),
    (
        "authority_impersonation",
        re.compile(r"\b(?:system|developer) message\s*:", re.I),
    ),
    (
        "tool_coercion",
        re.compile(r"\b(?:execute|run)\s+(?:this\s+)?(?:command|tool)\b", re.I),
    ),
)
_CONSTRUCTION_KEY = object()


class DocumentIngestionV1Error(RuntimeError):
    pass


class DocumentIngestionV1ContractError(ValueError):
    pass


class DocumentIngestionV1Denied(PermissionError):
    pass


@dataclass(frozen=True, slots=True)
class DocumentIngestionFeatureGateV1:
    enabled: bool

    def __post_init__(self) -> None:
        if type(self.enabled) is not bool:
            raise DocumentIngestionV1ContractError("feature gate must be exact bool")

    @classmethod
    def from_environ(
        cls, environ: dict[str, str] | os._Environ[str] | None = None
    ) -> "DocumentIngestionFeatureGateV1":
        source = os.environ if environ is None else environ
        return cls(source.get(FEATURE_FLAG) == ENABLED_VALUE)


def _id(name: str, value: object) -> str:
    if type(value) is not str or not _SAFE_ID.fullmatch(value):
        raise DocumentIngestionV1ContractError(f"{name} is invalid")
    return value


def _timestamp(name: str, value: object) -> int:
    if type(value) is not int or value < 0:
        raise DocumentIngestionV1ContractError(f"{name} is invalid")
    return value


def _canonical(value: object) -> bytes:
    try:
        return json.dumps(
            value,
            sort_keys=True,
            separators=(",", ":"),
            ensure_ascii=False,
            allow_nan=False,
        ).encode("utf-8")
    except (TypeError, ValueError) as exc:
        raise DocumentIngestionV1ContractError("value is not canonical JSON") from exc


def _mac(payload: object, key: bytes) -> str:
    return hmac.new(key, _canonical(payload), hashlib.sha256).hexdigest()


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _verify_entries(project_root: Path | str) -> None:
    root = Path(project_root).resolve()
    for relative, expected in (*SOURCE_ENTRY_ROOTS, *LAYERED_MEMORY_ENTRY_ROOTS):
        path = (root / relative).resolve()
        try:
            path.relative_to(root)
            actual = _sha(path.read_bytes())
        except (OSError, ValueError) as exc:
            raise DocumentIngestionV1Denied(
                "accepted entry evidence unavailable"
            ) from exc
        if not hmac.compare_digest(actual, expected):
            raise DocumentIngestionV1Denied("accepted entry evidence drift")


@dataclass(frozen=True, slots=True)
class DocumentEnvelopeV1:
    source_name: str
    logical_document_id: str
    revision_id: str
    filename: str
    media_type: str
    content_bytes: bytes
    observed_locator: str
    retrieved_at_ms: int

    def __post_init__(self) -> None:
        _id("source_name", self.source_name)
        _id("logical_document_id", self.logical_document_id)
        _id("revision_id", self.revision_id)
        if type(self.filename) is not str or not _FILENAME.fullmatch(self.filename):
            raise DocumentIngestionV1ContractError("filename is invalid")
        if self.media_type not in FORMATS:
            raise DocumentIngestionV1ContractError("media_type is unsupported")
        expected_suffixes = FORMATS[self.media_type][1]
        if Path(self.filename).suffix.casefold() not in expected_suffixes:
            raise DocumentIngestionV1ContractError("filename/media_type mismatch")
        if (
            type(self.content_bytes) is not bytes
            or not self.content_bytes
            or len(self.content_bytes) > MAX_DOCUMENT_BYTES
        ):
            raise DocumentIngestionV1ContractError("document bytes are invalid")
        if (
            type(self.observed_locator) is not str
            or not self.observed_locator
            or self.observed_locator != self.observed_locator.strip()
            or len(self.observed_locator) > 2_048
        ):
            raise DocumentIngestionV1ContractError("observed_locator is invalid")
        _timestamp("retrieved_at_ms", self.retrieved_at_ms)


@dataclass(frozen=True, slots=True)
class OcrBlockV1:
    text: str
    bbox: tuple[int, int, int, int]
    confidence_bp: int

    def __post_init__(self) -> None:
        if (
            type(self.text) is not str
            or not self.text.strip()
            or len(self.text) > MAX_UNIT_CHARS
            or type(self.bbox) is not tuple
            or len(self.bbox) != 4
            or any(type(value) is not int or value < 0 for value in self.bbox)
            or self.bbox[2] <= self.bbox[0]
            or self.bbox[3] <= self.bbox[1]
            or type(self.confidence_bp) is not int
            or not 0 <= self.confidence_bp <= 10_000
        ):
            raise DocumentIngestionV1ContractError("OCR block is invalid")


@dataclass(frozen=True, slots=True)
class OcrObservationV1:
    document_sha256: str
    engine: str
    engine_version: str
    blocks: tuple[OcrBlockV1, ...]

    def __post_init__(self) -> None:
        if not _HEX64.fullmatch(self.document_sha256):
            raise DocumentIngestionV1ContractError("OCR document hash is invalid")
        _id("OCR engine", self.engine)
        _id("OCR engine_version", self.engine_version)
        if (
            type(self.blocks) is not tuple
            or not self.blocks
            or len(self.blocks) > MAX_OCR_BLOCKS
            or any(type(block) is not OcrBlockV1 for block in self.blocks)
        ):
            raise DocumentIngestionV1ContractError("OCR blocks are invalid")


@dataclass(frozen=True, slots=True)
class RenderSurfaceV1:
    surface: str
    width_px: int
    height_px: int
    pixel_sha256: str
    issues: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        _id("render surface", self.surface)
        if (
            type(self.width_px) is not int
            or self.width_px <= 0
            or type(self.height_px) is not int
            or self.height_px <= 0
            or not _HEX64.fullmatch(self.pixel_sha256)
            or type(self.issues) is not tuple
            or any(type(item) is not str or not item for item in self.issues)
        ):
            raise DocumentIngestionV1ContractError("render surface is invalid")


@dataclass(frozen=True, slots=True)
class RenderObservationV1:
    document_sha256: str
    renderer: str
    renderer_version: str
    surfaces: tuple[RenderSurfaceV1, ...]

    def __post_init__(self) -> None:
        if not _HEX64.fullmatch(self.document_sha256):
            raise DocumentIngestionV1ContractError("render document hash is invalid")
        _id("renderer", self.renderer)
        _id("renderer_version", self.renderer_version)
        if (
            type(self.surfaces) is not tuple
            or not self.surfaces
            or len({surface.surface for surface in self.surfaces}) != len(self.surfaces)
            or any(type(surface) is not RenderSurfaceV1 for surface in self.surfaces)
        ):
            raise DocumentIngestionV1ContractError("render surfaces are invalid")


@dataclass(frozen=True, slots=True)
class CitationV1:
    citation_id: str
    source_id: str
    document_sha256: str
    revision_id: str
    unit_kind: str
    unit_index: int
    location: tuple[tuple[str, str], ...]
    locator: str
    text: str
    text_sha256: str
    content_trust: str = "untrusted_data"
    instructions_authority: bool = False


@dataclass(frozen=True, slots=True)
class ExtractedStatementV1:
    statement_id: str
    statement_kind: str
    text: str
    citation_ids: tuple[str, ...]
    extraction_method: str
    confidence_bp: int
    content_trust: str = "untrusted_data"
    instructions_authority: bool = False


@dataclass(frozen=True, slots=True)
class QaFindingV1:
    severity: str
    code: str
    surface: str
    detail: str


@dataclass(frozen=True, slots=True)
class IngestionResultV1:
    workspace_id: str
    principal_id: str
    source_name: str
    source_id: str
    logical_document_id: str
    revision_id: str
    filename: str
    media_type: str
    format: str
    document_sha256: str
    retrieved_at_ms: int
    source_verified_at_ms: int
    citations: tuple[CitationV1, ...]
    requirements: tuple[ExtractedStatementV1, ...]
    decisions: tuple[ExtractedStatementV1, ...]
    memory_candidates: tuple[LayeredMemorySpecV1, ...]
    qa_findings: tuple[QaFindingV1, ...]
    render_surfaces: tuple[RenderSurfaceV1, ...]
    render_status: str
    ocr_status: str
    poison_signals: tuple[str, ...]
    content_trust: str
    instructions_authority: bool
    result_sha256: str
    hmac_sha256: str


@dataclass(frozen=True, slots=True)
class VersionChangeV1:
    change_kind: str
    location: tuple[tuple[str, str], ...]
    before_citation_id: str | None
    after_citation_id: str | None
    similarity_bp: int


@dataclass(frozen=True, slots=True)
class VersionComparisonV1:
    logical_document_id: str
    before_revision_id: str
    after_revision_id: str
    changes: tuple[VersionChangeV1, ...]
    comparison_sha256: str
    content_trust: str = "untrusted_data"
    instructions_authority: bool = False


def _zip_guard(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ARCHIVE_ENTRIES:
                raise DocumentIngestionV1Denied("archive entry bound exceeded")
            names: set[str] = set()
            total = 0
            for info in infos:
                path = PurePosixPath(info.filename)
                normalized = info.filename.casefold()
                if (
                    normalized in names
                    or path.is_absolute()
                    or ".." in path.parts
                    or info.flag_bits & 0x1
                ):
                    raise DocumentIngestionV1Denied("unsafe archive entry")
                names.add(normalized)
                total += info.file_size
                if total > MAX_EXPANDED_BYTES or (
                    info.file_size > 1_000_000
                    and info.compress_size > 0
                    and info.file_size / info.compress_size > 200
                ):
                    raise DocumentIngestionV1Denied("archive expansion bound exceeded")
                if info.filename.endswith((".xml", ".rels")):
                    sample = archive.read(info)
                    if b"<!DOCTYPE" in sample.upper() or b"<!ENTITY" in sample.upper():
                        raise DocumentIngestionV1Denied("active XML declaration denied")
    except (zipfile.BadZipFile, RuntimeError) as exc:
        raise DocumentIngestionV1Denied("invalid document container") from exc


def _decode_text(data: bytes) -> str:
    try:
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            return data.decode("utf-16")
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentIngestionV1Denied(
            "text encoding must be UTF-8 or BOM UTF-16"
        ) from exc


def _unit(
    kind: str, location: dict[str, object], text: object
) -> tuple[str, dict[str, str], str] | None:
    value = "" if text is None else str(text)
    value = value.replace("\x00", "").strip()
    if not value:
        return None
    if len(value) > MAX_UNIT_CHARS:
        value = value[:MAX_UNIT_CHARS]
    return kind, {key: str(item) for key, item in location.items()}, value


def _extract_pdf(
    data: bytes,
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1]]:
    try:
        import pypdf
    except ImportError as exc:
        raise DocumentIngestionV1Error("pypdf is required") from exc
    units: list[tuple[str, dict[str, str], str]] = []
    qa: list[QaFindingV1] = []
    try:
        reader = pypdf.PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentIngestionV1Denied("encrypted PDF denied")
        for page_index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            if not lines:
                qa.append(
                    QaFindingV1(
                        "warning",
                        "blank_or_scanned_page",
                        f"page:{page_index}",
                        "No extractable text",
                    )
                )
            for line_index, line in enumerate(lines, start=1):
                item = _unit("pdf_line", {"page": page_index, "line": line_index}, line)
                if item:
                    units.append(item)
    except DocumentIngestionV1Denied:
        raise
    except Exception as exc:
        raise DocumentIngestionV1Denied("PDF parse failed") from exc
    return units, qa


def _extract_docx(
    data: bytes,
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1]]:
    _zip_guard(data)
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentIngestionV1Error("python-docx is required") from exc
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentIngestionV1Denied("DOCX parse failed") from exc
    units: list[tuple[str, dict[str, str], str]] = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        item = _unit("docx_paragraph", {"paragraph": index}, paragraph.text)
        if item:
            units.append(item)
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                item = _unit(
                    "docx_table_cell",
                    {"table": table_index, "row": row_index, "column": column_index},
                    cell.text,
                )
                if item:
                    units.append(item)
    return units, []


def _extract_lines(
    data: bytes, kind: str
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1]]:
    text = _decode_text(data)
    units = []
    for index, line in enumerate(text.splitlines(), start=1):
        item = _unit(f"{kind}_line", {"line": index}, line)
        if item:
            units.append(item)
    qa: list[QaFindingV1] = []
    if kind == "code":
        try:
            ast.parse(text) if text.strip() else None
        except SyntaxError as exc:
            qa.append(
                QaFindingV1(
                    "warning", "code_syntax_error", f"line:{exc.lineno or 0}", exc.msg
                )
            )
    return units, qa


def _extract_spreadsheet(
    data: bytes, media_type: str
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1]]:
    units: list[tuple[str, dict[str, str], str]] = []
    if media_type in {"text/csv", "text/tab-separated-values"}:
        delimiter = "," if media_type == "text/csv" else "\t"
        rows = csv.reader(io.StringIO(_decode_text(data)), delimiter=delimiter)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                item = _unit(
                    "spreadsheet_cell",
                    {"sheet": "Sheet1", "row": row_index, "column": column_index},
                    value,
                )
                if item:
                    units.append(item)
        return units, []
    if media_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }:
        _zip_guard(data)
        try:
            import openpyxl  # type: ignore[import-untyped]

            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        item = _unit(
                            "spreadsheet_cell",
                            {"sheet": sheet.title, "cell": cell.coordinate},
                            cell.value,
                        )
                        if item:
                            units.append(item)
            workbook.close()
            return units, []
        except Exception as exc:
            raise DocumentIngestionV1Denied("XLSX parse failed") from exc
    try:
        import pandas as pd  # type: ignore[import-untyped]

        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
        for sheet_name, frame in sheets.items():
            for row_index, row in frame.iterrows():
                for column_index, value in enumerate(row.tolist(), start=1):
                    if pd.isna(value):
                        continue
                    item = _unit(
                        "spreadsheet_cell",
                        {
                            "sheet": sheet_name,
                            "row": row_index + 1,
                            "column": column_index,
                        },
                        value,
                    )
                    if item:
                        units.append(item)
        return units, []
    except Exception as exc:
        raise DocumentIngestionV1Denied("spreadsheet parse failed") from exc


def _extract_presentation(
    data: bytes,
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1]]:
    _zip_guard(data)
    try:
        from pptx import Presentation

        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise DocumentIngestionV1Denied("PPTX parse failed") from exc
    units: list[tuple[str, dict[str, str], str]] = []
    qa: list[QaFindingV1] = []
    for slide_index, slide in enumerate(deck.slides, start=1):
        count = 0
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "has_text_frame", False):
                item = _unit(
                    "presentation_shape",
                    {"slide": slide_index, "shape": shape_index},
                    shape.text,
                )
                if item:
                    units.append(item)
                    count += 1
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows, start=1):
                    for column_index, cell in enumerate(row.cells, start=1):
                        item = _unit(
                            "presentation_table_cell",
                            {
                                "slide": slide_index,
                                "shape": shape_index,
                                "row": row_index,
                                "column": column_index,
                            },
                            cell.text,
                        )
                        if item:
                            units.append(item)
                            count += 1
        if count == 0:
            qa.append(
                QaFindingV1(
                    "warning",
                    "slide_without_extractable_text",
                    f"slide:{slide_index}",
                    "No text found",
                )
            )
    return units, qa


def _native_visual_qa(
    data: bytes, format_name: str
) -> tuple[str, list[QaFindingV1], tuple[RenderSurfaceV1, ...]]:
    if format_name not in {"pdf", "image"}:
        return "not_run", [], ()
    surfaces: list[RenderSurfaceV1] = []
    qa: list[QaFindingV1] = []
    if format_name == "image":
        try:
            from PIL import Image, ImageStat

            decoded = Image.open(io.BytesIO(data))
            decoded.verify()
            image = Image.open(io.BytesIO(data)).convert("RGB")
            stat = ImageStat.Stat(image.convert("L"))
            if stat.var[0] < 1.0:
                qa.append(
                    QaFindingV1(
                        "warning",
                        "low_visual_variance",
                        "image:1",
                        "Image is nearly uniform",
                    )
                )
            surfaces.append(
                RenderSurfaceV1(
                    "image:1", image.width, image.height, _sha(image.tobytes())
                )
            )
        except Exception as exc:
            raise DocumentIngestionV1Denied("image decode/render failed") from exc
        return "native_rendered", qa, tuple(surfaces)
    try:
        import pypdfium2 as pdfium  # type: ignore[import-untyped]

        document = pdfium.PdfDocument(data)
        for page_index in range(len(document)):
            bitmap = document[page_index].render(scale=0.5)
            image = bitmap.to_pil().convert("RGB")
            surfaces.append(
                RenderSurfaceV1(
                    f"page:{page_index + 1}",
                    image.width,
                    image.height,
                    _sha(image.tobytes()),
                )
            )
            bitmap.close()
        document.close()
    except ImportError:
        qa.append(
            QaFindingV1(
                "warning",
                "render_engine_unavailable",
                "document",
                "pypdfium2 unavailable",
            )
        )
        return "not_run", qa, ()
    except Exception as exc:
        raise DocumentIngestionV1Denied("PDF render failed") from exc
    return "native_rendered", qa, tuple(surfaces)


def _extract_image(
    data: bytes, ocr: OcrObservationV1 | None
) -> tuple[list[tuple[str, dict[str, str], str]], list[QaFindingV1], str]:
    digest = _sha(data)
    if ocr is None:
        return (
            [],
            [QaFindingV1("info", "ocr_not_supplied", "image:1", "No OCR observation")],
            "not_run",
        )
    if type(ocr) is not OcrObservationV1 or ocr.document_sha256 != digest:
        raise DocumentIngestionV1Denied("OCR observation binding denied")
    units = []
    for index, block in enumerate(ocr.blocks, start=1):
        item = _unit(
            "ocr_block",
            {
                "image": 1,
                "block": index,
                "bbox": ",".join(str(value) for value in block.bbox),
                "confidence_bp": block.confidence_bp,
                "engine": f"{ocr.engine}@{ocr.engine_version}",
            },
            block.text,
        )
        if item:
            units.append(item)
    return units, [], "observed"


def _bounded(
    units: list[tuple[str, dict[str, str], str]],
) -> list[tuple[str, dict[str, str], str]]:
    if len(units) > MAX_UNITS:
        raise DocumentIngestionV1Denied("document unit bound exceeded")
    total = sum(len(text) for _, _, text in units)
    if total > MAX_TOTAL_TEXT_CHARS:
        raise DocumentIngestionV1Denied("document text bound exceeded")
    return units


def _citation(
    source: ApprovedSourceRecordV1,
    envelope: DocumentEnvelopeV1,
    document_sha256: str,
    index: int,
    raw: tuple[str, dict[str, str], str],
) -> CitationV1:
    kind, location_map, text = raw
    location = tuple(sorted(location_map.items()))
    text_sha256 = _sha(text.encode("utf-8"))
    material = {
        "source_id": source.source_id,
        "document_sha256": document_sha256,
        "revision_id": envelope.revision_id,
        "unit_kind": kind,
        "unit_index": index,
        "location": location,
        "text_sha256": text_sha256,
    }
    identifier = f"cite_{_sha(_canonical(material))[:32]}"
    fragment = "&".join(f"{key}={value}" for key, value in location)
    return CitationV1(
        identifier,
        source.source_id,
        document_sha256,
        envelope.revision_id,
        kind,
        index,
        location,
        f"{source.citation}#{fragment}",
        text,
        text_sha256,
    )


def _statements(
    citations: tuple[CitationV1, ...], pattern: re.Pattern[str], kind: str
) -> tuple[ExtractedStatementV1, ...]:
    values = []
    for citation in citations:
        if not pattern.search(citation.text):
            continue
        payload = {
            "kind": kind,
            "text_sha256": citation.text_sha256,
            "citation_id": citation.citation_id,
        }
        values.append(
            ExtractedStatementV1(
                f"stmt_{_sha(_canonical(payload))[:32]}",
                kind,
                citation.text,
                (citation.citation_id,),
                "deterministic_pattern_v1",
                7_500,
            )
        )
    return tuple(values)


def _result_payload(result: IngestionResultV1) -> dict[str, object]:
    payload = asdict(result)
    payload.pop("result_sha256", None)
    payload.pop("hmac_sha256", None)
    return payload


class GovernedDocumentIngestorV1:
    __slots__ = (
        "_sources",
        "_workspace_id",
        "_principal_id",
        "_integrity_key",
        "_key_fingerprint",
    )

    def __init__(
        self,
        *,
        construction_key: object,
        sources: ApprovedSourceRegistryV1,
        integrity_key: bytes,
    ) -> None:
        if construction_key is not _CONSTRUCTION_KEY:
            raise DocumentIngestionV1ContractError(
                "use create_governed_document_ingestor_v1"
            )
        if type(sources) is not ApprovedSourceRegistryV1:
            raise DocumentIngestionV1ContractError(
                "exact ApprovedSourceRegistryV1 required"
            )
        if type(integrity_key) is not bytes or not 32 <= len(integrity_key) <= 64:
            raise DocumentIngestionV1ContractError("integrity_key is invalid")
        fingerprint = hashlib.sha256(integrity_key).hexdigest()
        if sources.key_fingerprint_sha256 != fingerprint:
            raise DocumentIngestionV1Denied("source registry key binding denied")
        self._sources = sources
        self._workspace_id = sources.workspace_id
        self._principal_id = sources.principal_id
        self._integrity_key = integrity_key
        self._key_fingerprint = fingerprint

    def __init_subclass__(cls, **_kwargs: object) -> None:
        raise TypeError("GovernedDocumentIngestorV1 cannot be subclassed")

    def __reduce__(self) -> NoReturn:
        raise TypeError("GovernedDocumentIngestorV1 cannot be serialized")

    def _attest(self) -> None:
        if (
            type(self._sources) is not ApprovedSourceRegistryV1
            or self._sources.workspace_id != self._workspace_id
            or self._sources.principal_id != self._principal_id
            or self._sources.key_fingerprint_sha256 != self._key_fingerprint
            or hashlib.sha256(self._integrity_key).hexdigest() != self._key_fingerprint
        ):
            raise DocumentIngestionV1Denied("ingestor binding drift denied")

    def ingest(
        self,
        envelope: DocumentEnvelopeV1,
        *,
        now_ms: int,
        ocr_observation: OcrObservationV1 | None = None,
        render_observation: RenderObservationV1 | None = None,
    ) -> IngestionResultV1:
        if type(envelope) is not DocumentEnvelopeV1:
            raise DocumentIngestionV1ContractError("exact document envelope required")
        _timestamp("now_ms", now_ms)
        self._attest()
        try:
            source = self._sources.get(
                envelope.source_name,
                now_ms=now_ms,
                require_fresh=True,
            )
        except ApprovedSourceV1Denied as exc:
            raise DocumentIngestionV1Denied("approved source unavailable") from exc
        if (
            source.workspace_id != self._workspace_id
            or source.principal_id != self._principal_id
            or source.locator != envelope.observed_locator
            or envelope.retrieved_at_ms > now_ms
        ):
            raise DocumentIngestionV1Denied("document/source binding denied")
        document_sha = _sha(envelope.content_bytes)
        if source.locator_kind == "artifact_alias" and (
            source.artifact_sha256 != document_sha
        ):
            raise DocumentIngestionV1Denied("artifact bytes hash drift")
        format_name = FORMATS[envelope.media_type][0]
        qa: list[QaFindingV1] = []
        ocr_status = "not_applicable"
        if format_name == "pdf":
            units, parsed_qa = _extract_pdf(envelope.content_bytes)
        elif format_name == "docx":
            units, parsed_qa = _extract_docx(envelope.content_bytes)
        elif format_name in {"markdown", "text", "code"}:
            units, parsed_qa = _extract_lines(envelope.content_bytes, format_name)
        elif format_name == "spreadsheet":
            units, parsed_qa = _extract_spreadsheet(
                envelope.content_bytes, envelope.media_type
            )
        elif format_name == "presentation":
            units, parsed_qa = _extract_presentation(envelope.content_bytes)
        elif format_name == "image":
            units, parsed_qa, ocr_status = _extract_image(
                envelope.content_bytes, ocr_observation
            )
        else:
            raise DocumentIngestionV1ContractError("format dispatch drift")
        qa.extend(parsed_qa)
        units = _bounded(units)
        citations = tuple(
            _citation(source, envelope, document_sha, index, raw)
            for index, raw in enumerate(units, start=1)
        )
        requirements = _statements(citations, _REQ, "requirement")
        decisions = _statements(citations, _DECISION, "decision")
        citation_ids = {citation.citation_id for citation in citations}
        if any(
            identifier not in citation_ids
            for statement in (*requirements, *decisions)
            for identifier in statement.citation_ids
        ):
            raise DocumentIngestionV1Denied("statement grounding drift")
        render_status, native_qa, native_surfaces = _native_visual_qa(
            envelope.content_bytes, format_name
        )
        qa.extend(native_qa)
        surfaces = native_surfaces
        if render_observation is not None:
            if (
                type(render_observation) is not RenderObservationV1
                or render_observation.document_sha256 != document_sha
            ):
                raise DocumentIngestionV1Denied("render observation binding denied")
            render_status = "native_and_observed" if native_surfaces else "observed"
            surfaces = (*native_surfaces, *render_observation.surfaces)
        for surface in surfaces:
            for issue in surface.issues:
                qa.append(
                    QaFindingV1(
                        "warning", "renderer_reported_issue", surface.surface, issue
                    )
                )
        poison = tuple(
            sorted(
                {
                    name
                    for citation in citations
                    for name, pattern in _POISON
                    if pattern.search(citation.text)
                }
            )
        )
        candidates: list[LayeredMemorySpecV1] = []
        for statement in (*requirements, *decisions):
            candidates.append(
                LayeredMemorySpecV1(
                    layer=(
                        "decision"
                        if statement.statement_kind == "decision"
                        else "semantic_institutional"
                    ),
                    entity_kind=statement.statement_kind,
                    entity_key=statement.statement_id,
                    content=statement.text,
                    source_ids=statement.citation_ids,
                    sensitivity=source.sensitivity,
                    confidence_bp=min(statement.confidence_bp, source.credibility_bp),
                    valid_from_ms=source.valid_from_ms,
                    valid_until_ms=source.valid_until_ms,
                    fresh_until_ms=source.fresh_until_ms,
                    retention_until_ms=source.valid_until_ms,
                    tags=("document-ingestion", statement.statement_kind),
                    status="candidate",
                )
            )
        provisional = IngestionResultV1(
            self._workspace_id,
            self._principal_id,
            source.source_name,
            source.source_id,
            envelope.logical_document_id,
            envelope.revision_id,
            envelope.filename,
            envelope.media_type,
            format_name,
            document_sha,
            envelope.retrieved_at_ms,
            now_ms,
            citations,
            requirements,
            decisions,
            tuple(candidates),
            tuple(qa),
            tuple(surfaces),
            render_status,
            ocr_status,
            poison,
            "untrusted_data",
            False,
            "",
            "",
        )
        payload = _result_payload(provisional)
        result_sha = _sha(_canonical(payload))
        result = IngestionResultV1(
            **{
                **asdict(provisional),
                "citations": citations,
                "requirements": requirements,
                "decisions": decisions,
                "memory_candidates": tuple(candidates),
                "qa_findings": tuple(qa),
                "render_surfaces": tuple(surfaces),
                "result_sha256": result_sha,
                "hmac_sha256": _mac(
                    {"result_sha256": result_sha, "payload": payload},
                    self._integrity_key,
                ),
            }
        )
        self._attest()
        return result

    def verify_result(self, result: IngestionResultV1, *, now_ms: int) -> None:
        if type(result) is not IngestionResultV1:
            raise DocumentIngestionV1ContractError("exact ingestion result required")
        _timestamp("now_ms", now_ms)
        self._attest()
        payload = _result_payload(result)
        expected_sha = _sha(_canonical(payload))
        if (
            result.workspace_id != self._workspace_id
            or result.principal_id != self._principal_id
            or not hmac.compare_digest(result.result_sha256, expected_sha)
            or not hmac.compare_digest(
                result.hmac_sha256,
                _mac(
                    {"result_sha256": expected_sha, "payload": payload},
                    self._integrity_key,
                ),
            )
        ):
            raise DocumentIngestionV1Denied("ingestion result integrity drift")
        try:
            source = self._sources.get(
                result.source_name, now_ms=now_ms, require_fresh=True
            )
        except ApprovedSourceV1Denied as exc:
            raise DocumentIngestionV1Denied(
                "ingestion source no longer available"
            ) from exc
        if source.source_id != result.source_id:
            raise DocumentIngestionV1Denied("ingestion source identity drift")

    def compare_versions(
        self,
        before: IngestionResultV1,
        after: IngestionResultV1,
        *,
        now_ms: int,
    ) -> VersionComparisonV1:
        self.verify_result(before, now_ms=now_ms)
        self.verify_result(after, now_ms=now_ms)
        if (
            before.logical_document_id != after.logical_document_id
            or before.source_id != after.source_id
            or before.revision_id == after.revision_id
        ):
            raise DocumentIngestionV1ContractError(
                "version comparison identity mismatch"
            )
        prior = {(item.unit_kind, item.location): item for item in before.citations}
        current = {(item.unit_kind, item.location): item for item in after.citations}
        changes: list[VersionChangeV1] = []
        for key in sorted(set(prior) | set(current)):
            old = prior.get(key)
            new = current.get(key)
            if old is None:
                assert new is not None
                changes.append(
                    VersionChangeV1("added", key[1], None, new.citation_id, 0)
                )
            elif new is None:
                changes.append(
                    VersionChangeV1("removed", key[1], old.citation_id, None, 0)
                )
            elif old.text_sha256 != new.text_sha256:
                similarity = int(
                    difflib.SequenceMatcher(None, old.text, new.text).ratio() * 10_000
                )
                changes.append(
                    VersionChangeV1(
                        "modified", key[1], old.citation_id, new.citation_id, similarity
                    )
                )
        payload = {
            "logical_document_id": before.logical_document_id,
            "before_revision_id": before.revision_id,
            "after_revision_id": after.revision_id,
            "changes": [asdict(change) for change in changes],
        }
        return VersionComparisonV1(
            before.logical_document_id,
            before.revision_id,
            after.revision_id,
            tuple(changes),
            _sha(_canonical(payload)),
        )


def create_governed_document_ingestor_v1(
    *,
    gate: DocumentIngestionFeatureGateV1 | None = None,
    sources: ApprovedSourceRegistryV1 | None = None,
    integrity_key: bytes | None = None,
    project_root: Path | str | None = None,
) -> GovernedDocumentIngestorV1 | None:
    selected = DocumentIngestionFeatureGateV1.from_environ() if gate is None else gate
    if type(selected) is not DocumentIngestionFeatureGateV1:
        raise DocumentIngestionV1ContractError("sealed feature gate required")
    if not selected.enabled:
        return None
    _verify_entries(
        Path(__file__).resolve().parents[1] if project_root is None else project_root
    )
    if sources is None or integrity_key is None:
        raise DocumentIngestionV1ContractError(
            "enabled ingestor requires complete host bindings"
        )
    return GovernedDocumentIngestorV1(
        construction_key=_CONSTRUCTION_KEY,
        sources=sources,
        integrity_key=integrity_key,
    )


__all__ = [
    "CitationV1",
    "DocumentEnvelopeV1",
    "DocumentIngestionFeatureGateV1",
    "DocumentIngestionV1ContractError",
    "DocumentIngestionV1Denied",
    "DocumentIngestionV1Error",
    "ExtractedStatementV1",
    "FEATURE_FLAG",
    "FORMATS",
    "GovernedDocumentIngestorV1",
    "IngestionResultV1",
    "OcrBlockV1",
    "OcrObservationV1",
    "QaFindingV1",
    "RenderObservationV1",
    "RenderSurfaceV1",
    "SCHEMA",
    "VersionChangeV1",
    "VersionComparisonV1",
    "create_governed_document_ingestor_v1",
]
